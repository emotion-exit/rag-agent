import io
import os
import re

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.services.document_processors.common import (
    Block,
    build_anchor_text,
    guess_content_type,
    find_meaningful_paragraph,
    infer_heading_level,
    looks_like_heading,
    normalize_block,
    update_heading_stack,
    build_section_payload,
)
from app.services.document_processors.text_processor import extract_text_sections


def extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def extract_docx_sections(file_bytes: bytes) -> list[dict]:
    doc = Document(io.BytesIO(file_bytes))
    sections: list[dict] = []
    heading_stack: list[str] = []
    current_blocks: list[Block] = []
    current_title = ""
    paragraph_index = 0
    block_index = 0

    def flush_current_section() -> None:
        nonlocal current_blocks
        if not current_blocks:
            return

        sections.append(
            build_section_payload(
                blocks=current_blocks,
                section_title=current_title,
                heading_path=" > ".join(heading_stack),
                source_label=current_title or "正文文本",
                source_page=0,
            )
        )
        current_blocks = []

    for paragraph in doc.paragraphs:
        text = normalize_block(paragraph.text)
        if not text:
            continue

        paragraph_index += 1
        block_index += 1

        if is_docx_heading(paragraph):
            flush_current_section()
            level = get_docx_heading_level(paragraph)
            heading_stack = update_heading_stack(heading_stack, text, level)
            current_title = text
            continue

        current_blocks.append(
            {
                "text": text,
                "paragraph_index": paragraph_index,
                "block_index": block_index,
            }
        )

    flush_current_section()

    if sections:
        return sections

    full_text = "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    return extract_text_sections(full_text, source_label="正文文本")


def extract_docx_images(file_bytes: bytes) -> list[dict]:
    doc = Document(io.BytesIO(file_bytes))
    sequence: list[dict] = []
    heading_stack: list[str] = []
    current_title = ""
    paragraph_index = 0
    image_counter = 0

    for child in doc.element.body.iterchildren():
        if not child.tag.endswith("}p"):
            continue

        paragraph = Paragraph(child, doc)
        text = normalize_block(paragraph.text)
        is_heading = bool(text) and is_docx_heading(paragraph)

        if text:
            paragraph_index += 1
            if is_heading:
                level = get_docx_heading_level(paragraph)
                heading_stack = update_heading_stack(heading_stack, text, level)
                current_title = text

            sequence.append(
                {
                    "type": "paragraph",
                    "text": text,
                    "paragraph_index": paragraph_index,
                    "heading_path": " > ".join(heading_stack),
                    "section_title": current_title,
                    "is_heading": is_heading,
                }
            )

        for image in extract_paragraph_images(paragraph):
            image_counter += 1
            sequence.append(
                {
                    "type": "image",
                    "bytes": image["bytes"],
                    "filename": image["filename"],
                    "content_type": image["content_type"],
                    "source_label": f"文档图片 {image_counter}",
                    "heading_path": " > ".join(heading_stack),
                    "section_title": current_title,
                }
            )

    images: list[dict] = []
    for index, item in enumerate(sequence):
        if item.get("type") != "image":
            continue

        anchor_before = find_meaningful_paragraph(sequence, index, direction=-1)
        anchor_after = find_meaningful_paragraph(sequence, index, direction=1)
        anchor_paragraph_index = 0
        if anchor_before is not None:
            anchor_paragraph_index = int(anchor_before.get("paragraph_index", 0) or 0)
        elif anchor_after is not None:
            anchor_paragraph_index = int(anchor_after.get("paragraph_index", 0) or 0)

        images.append(
            {
                "filename": item["filename"],
                "bytes": item["bytes"],
                "content_type": item["content_type"],
                "source_label": item["source_label"],
                "source_page": 0,
                "paragraph_index": anchor_paragraph_index,
                "anchor_before": str(anchor_before.get("text", "") if anchor_before else ""),
                "anchor_after": str(anchor_after.get("text", "") if anchor_after else ""),
                "anchor_text": build_anchor_text(anchor_before, anchor_after),
                "heading_path": str(
                    item.get("heading_path")
                    or (anchor_before or {}).get("heading_path", "")
                    or (anchor_after or {}).get("heading_path", "")
                ),
                "section_title": str(
                    item.get("section_title")
                    or (anchor_before or {}).get("section_title", "")
                    or (anchor_after or {}).get("section_title", "")
                ),
            }
        )

    return images


def extract_paragraph_images(paragraph: Paragraph) -> list[dict]:
    images: list[dict] = []
    seen_rel_ids: set[str] = set()

    for blip in paragraph._element.xpath(".//*[local-name()='blip']"):
        rel_id = blip.get(qn("r:embed"))
        if not rel_id or rel_id in seen_rel_ids:
            continue
        seen_rel_ids.add(rel_id)

        image_part = paragraph.part.related_parts.get(rel_id)
        if image_part is None:
            continue

        part_name = os.path.basename(str(image_part.partname))
        images.append(
            {
                "filename": part_name,
                "bytes": image_part.blob,
                "content_type": guess_content_type(part_name),
            }
        )

    return images


def is_docx_heading(paragraph) -> bool:
    style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
    if looks_like_heading(paragraph.text):
        return True
    return bool(re.match(r"^(heading|标题)\s*\d*$", style_name.strip(), flags=re.IGNORECASE))


def get_docx_heading_level(paragraph) -> int:
    style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
    match = re.search(r"(\d+)", style_name)
    if match:
        return max(1, min(int(match.group(1)), 6))
    return infer_heading_level(paragraph.text)
