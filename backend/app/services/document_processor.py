import io
import os
import re
import zipfile
from importlib import import_module

import chardet

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


TEXT_SOURCE_TYPE = "text"
DEFAULT_CHUNK_SIZE = 240
DEFAULT_CHUNK_OVERLAP = 24
TITLE_MAX_LENGTH = 80
DOCX_IMAGE_PROXIMITY = 2
PDF_BLOCK_PROXIMITY = 2


Block = dict[str, int | str | tuple[float, float, float, float]]
Section = dict[str, object]
ImageInfo = dict[str, object]


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from supported file types."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    if ext == ".docx":
        return _extract_docx(file_bytes)
    if ext in (".txt", ".md", ".rst", ".csv"):
        return _extract_text(file_bytes)

    return _extract_text(file_bytes)


def extract_document_chunks(
    file_bytes: bytes,
    filename: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[dict]:
    sections = extract_document_sections(file_bytes, filename)
    chunks: list[dict] = []

    for section in sections:
        blocks = list(section.get("blocks", []))
        if not blocks:
            continue

        for chunk in _chunk_section_blocks(
            blocks,
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        ):
            chunks.append(
                {
                    "content": chunk["content"],
                    "source_type": TEXT_SOURCE_TYPE,
                    "source_label": str(section.get("source_label", "正文文本")),
                    "source_page": int(section.get("source_page", 0) or 0),
                    "section_title": str(section.get("section_title", "")).strip(),
                    "heading_path": str(section.get("heading_path", "")).strip(),
                    "paragraph_index_start": int(chunk.get("paragraph_index_start", 0) or 0),
                    "paragraph_index_end": int(chunk.get("paragraph_index_end", 0) or 0),
                    "block_index_start": int(chunk.get("block_index_start", 0) or 0),
                    "block_index_end": int(chunk.get("block_index_end", 0) or 0),
                }
            )

    return chunks


def extract_document_sections(file_bytes: bytes, filename: str) -> list[dict]:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".docx":
        return _extract_docx_sections(file_bytes)
    if ext == ".pdf":
        return _extract_pdf_sections(file_bytes)
    return _extract_text_sections(
        extract_text_from_file(file_bytes, filename),
        source_label="正文文本",
    )


def extract_document_images(file_bytes: bytes, filename: str) -> list[dict]:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".docx":
        return _extract_docx_images(file_bytes)
    if ext == ".pdf":
        return _extract_pdf_images(file_bytes)
    return []


def _extract_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    texts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text.strip())
    return "\n\n".join(texts)


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _extract_text(file_bytes: bytes) -> str:
    detected = chardet.detect(file_bytes)
    encoding = detected.get("encoding") or "utf-8"
    return file_bytes.decode(encoding, errors="replace")


def _extract_docx_sections(file_bytes: bytes) -> list[dict]:
    doc = Document(io.BytesIO(file_bytes))
    sections: list[dict] = []
    heading_stack: list[str] = []
    current_blocks: list[Block] = []
    current_title = ""
    paragraph_index = 0
    block_index = 0

    def flush_current_section() -> None:
        nonlocal current_blocks
        if not current_blocks:
            return

        sections.append(
            _build_section_payload(
                blocks=current_blocks,
                section_title=current_title,
                heading_path=" > ".join(heading_stack),
                source_label=current_title or "正文文本",
                source_page=0,
            )
        )
        current_blocks = []

    for paragraph in doc.paragraphs:
        text = _normalize_block(paragraph.text)
        if not text:
            continue

        paragraph_index += 1
        block_index += 1

        if _is_docx_heading(paragraph):
            flush_current_section()
            level = _get_docx_heading_level(paragraph)
            heading_stack = _update_heading_stack(heading_stack, text, level)
            current_title = text
            continue

        current_blocks.append(
            {
                "text": text,
                "paragraph_index": paragraph_index,
                "block_index": block_index,
            }
        )

    flush_current_section()

    if sections:
        return sections

    full_text = "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    return _extract_text_sections(full_text, source_label="正文文本")


def _extract_pdf_sections(file_bytes: bytes) -> list[dict]:
    fitz = _load_pymupdf()

    document = fitz.open(stream=file_bytes, filetype="pdf")
    sections: list[dict] = []

    try:
        for page_number in range(1, len(document) + 1):
            page = document.load_page(page_number - 1)
            page_blocks = _extract_pdf_text_blocks(page, page_number)
            if not page_blocks:
                continue

            sections.extend(
                _split_blocks_into_sections(
                    page_blocks,
                    source_label=f"第 {page_number} 页正文",
                    source_page=page_number,
                )
            )
    finally:
        document.close()

    return sections


def _extract_text_sections(
    text: str,
    source_label: str,
    source_page: int = 0,
) -> list[dict]:
    paragraphs = [
        _normalize_block(paragraph)
        for paragraph in re.split(r"\n\s*\n", text)
        if _normalize_block(paragraph)
    ]
    if not paragraphs:
        return []

    blocks: list[Block] = [
        {
            "text": paragraph,
            "paragraph_index": index,
            "block_index": index,
        }
        for index, paragraph in enumerate(paragraphs, start=1)
    ]

    return _split_blocks_into_sections(
        blocks,
        source_label=source_label,
        source_page=source_page,
    )


def _extract_docx_images(file_bytes: bytes) -> list[dict]:
    doc = Document(io.BytesIO(file_bytes))
    sequence: list[dict] = []
    heading_stack: list[str] = []
    current_title = ""
    paragraph_index = 0
    image_counter = 0

    for child in doc.element.body.iterchildren():
        if not child.tag.endswith("}p"):
            continue

        paragraph = Paragraph(child, doc)
        text = _normalize_block(paragraph.text)
        is_heading = bool(text) and _is_docx_heading(paragraph)

        if text:
            paragraph_index += 1
            if is_heading:
                level = _get_docx_heading_level(paragraph)
                heading_stack = _update_heading_stack(heading_stack, text, level)
                current_title = text

            sequence.append(
                {
                    "type": "paragraph",
                    "text": text,
                    "paragraph_index": paragraph_index,
                    "heading_path": " > ".join(heading_stack),
                    "section_title": current_title,
                    "is_heading": is_heading,
                }
            )

        for image in _extract_paragraph_images(paragraph):
            image_counter += 1
            sequence.append(
                {
                    "type": "image",
                    "bytes": image["bytes"],
                    "filename": image["filename"],
                    "content_type": image["content_type"],
                    "source_label": f"文档图片 {image_counter}",
                    "heading_path": " > ".join(heading_stack),
                    "section_title": current_title,
                }
            )

    images: list[dict] = []
    for index, item in enumerate(sequence):
        if item.get("type") != "image":
            continue

        anchor_before = _find_meaningful_paragraph(sequence, index, direction=-1)
        anchor_after = _find_meaningful_paragraph(sequence, index, direction=1)
        anchor_paragraph_index = 0
        if anchor_before is not None:
            anchor_paragraph_index = int(anchor_before.get("paragraph_index", 0) or 0)
        elif anchor_after is not None:
            anchor_paragraph_index = int(anchor_after.get("paragraph_index", 0) or 0)

        images.append(
            {
                "filename": item["filename"],
                "bytes": item["bytes"],
                "content_type": item["content_type"],
                "source_label": item["source_label"],
                "source_page": 0,
                "paragraph_index": anchor_paragraph_index,
                "anchor_before": str(anchor_before.get("text", "") if anchor_before else ""),
                "anchor_after": str(anchor_after.get("text", "") if anchor_after else ""),
                "anchor_text": _build_anchor_text(anchor_before, anchor_after),
                "heading_path": str(
                    item.get("heading_path")
                    or (anchor_before or {}).get("heading_path", "")
                    or (anchor_after or {}).get("heading_path", "")
                ),
                "section_title": str(
                    item.get("section_title")
                    or (anchor_before or {}).get("section_title", "")
                    or (anchor_after or {}).get("section_title", "")
                ),
            }
        )

    return images


def _extract_pdf_images(file_bytes: bytes) -> list[dict]:
    fitz = _load_pymupdf()

    document = fitz.open(stream=file_bytes, filetype="pdf")
    images: list[dict] = []

    try:
        for page_number in range(1, len(document) + 1):
            page = document.load_page(page_number - 1)
            text_blocks = _extract_pdf_text_blocks(page, page_number)
            seen_xrefs: set[int] = set()

            for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                xref = int(image_info[0])
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)

                image_data = document.extract_image(xref)
                image_bytes = image_data.get("image")
                if not image_bytes:
                    continue

                rects = page.get_image_rects(xref)
                bbox = rects[0] if rects else None
                anchor_before, anchor_after = _find_nearest_pdf_blocks(text_blocks, bbox)
                nearest_block = anchor_before or anchor_after
                ext = str(image_data.get("ext", "png")).lower()

                images.append(
                    {
                        "filename": f"page-{page_number}-image-{image_index}.{ext}",
                        "bytes": image_bytes,
                        "content_type": _guess_content_type(ext),
                        "source_label": f"第 {page_number} 页图片 {image_index}",
                        "source_page": page_number,
                        "block_index": int((nearest_block or {}).get("block_index", 0) or 0),
                        "anchor_before": str((anchor_before or {}).get("text", "")),
                        "anchor_after": str((anchor_after or {}).get("text", "")),
                        "anchor_text": _build_anchor_text(anchor_before, anchor_after),
                        "heading_path": str((nearest_block or {}).get("heading_path", "")),
                        "section_title": str((nearest_block or {}).get("section_title", "")),
                    }
                )
    finally:
        document.close()

    return images


def _extract_paragraph_images(paragraph: Paragraph) -> list[dict]:
    images: list[dict] = []
    seen_rel_ids: set[str] = set()

    for blip in paragraph._element.xpath(".//*[local-name()='blip']"):
        rel_id = blip.get(qn("r:embed"))
        if not rel_id or rel_id in seen_rel_ids:
            continue
        seen_rel_ids.add(rel_id)

        image_part = paragraph.part.related_parts.get(rel_id)
        if image_part is None:
            continue

        part_name = os.path.basename(str(image_part.partname))
        images.append(
            {
                "filename": part_name,
                "bytes": image_part.blob,
                "content_type": _guess_content_type(part_name),
            }
        )

    return images


def _extract_pdf_text_blocks(page, page_number: int) -> list[Block]:
    blocks: list[Block] = []
    heading_stack: list[str] = []
    current_title = ""
    block_index = 0

    for raw_block in page.get_text("blocks", sort=True):
        if len(raw_block) < 5:
            continue
        if len(raw_block) >= 7 and int(raw_block[6]) != 0:
            continue

        text = _normalize_block(str(raw_block[4]))
        if not text:
            continue

        block_index += 1
        if _looks_like_heading(text):
            level = _infer_heading_level(text)
            heading_stack = _update_heading_stack(heading_stack, text, level)
            current_title = text
            continue

        blocks.append(
            {
                "text": text,
                "paragraph_index": block_index,
                "block_index": block_index,
                "source_page": page_number,
                "bbox": (
                    float(raw_block[0]),
                    float(raw_block[1]),
                    float(raw_block[2]),
                    float(raw_block[3]),
                ),
                "heading_path": " > ".join(heading_stack),
                "section_title": current_title,
            }
        )

    return blocks


def _split_blocks_into_sections(
    blocks: list[Block],
    source_label: str,
    source_page: int,
) -> list[Section]:
    sections: list[Section] = []
    heading_stack: list[str] = []
    current_title = ""
    current_blocks: list[Block] = []

    def flush_current_section() -> None:
        nonlocal current_blocks
        if not current_blocks:
            return

        sections.append(
            _build_section_payload(
                blocks=current_blocks,
                section_title=current_title,
                heading_path=" > ".join(heading_stack),
                source_label=current_title or source_label,
                source_page=source_page,
            )
        )
        current_blocks = []

    for block in blocks:
        text = _normalize_block(str(block.get("text", "")))
        if _looks_like_heading(text):
            flush_current_section()
            level = _infer_heading_level(text)
            heading_stack = _update_heading_stack(heading_stack, text, level)
            current_title = text
            continue

        current_blocks.append(block)

    flush_current_section()

    return sections or [
        _build_section_payload(
            blocks=blocks,
            section_title="",
            heading_path="",
            source_label=source_label,
            source_page=source_page,
        )
    ]


def _build_section_payload(
    blocks: list[Block],
    section_title: str,
    heading_path: str,
    source_label: str,
    source_page: int,
) -> Section:
    return {
        "blocks": blocks,
        "content": _join_blocks([str(block.get("text", "")) for block in blocks]),
        "section_title": section_title,
        "heading_path": heading_path,
        "source_label": source_label,
        "source_page": source_page,
        "paragraph_index_start": _range_start(blocks, "paragraph_index"),
        "paragraph_index_end": _range_end(blocks, "paragraph_index"),
        "block_index_start": _range_start(blocks, "block_index"),
        "block_index_end": _range_end(blocks, "block_index"),
    }


def _chunk_section_blocks(
    blocks: list[Block],
    chunk_size: int,
    chunk_overlap: int,
) -> list[dict]:
    units: list[Block] = []
    for block in blocks:
        text = str(block.get("text", ""))
        if len(text) <= chunk_size:
            units.append(block)
        else:
            units.extend(_split_long_block(block, chunk_size))

    chunks: list[dict] = []
    current_units: list[Block] = []

    for unit in units:
        candidate_units = current_units + [unit]
        candidate_length = len(_join_blocks([str(item.get("text", "")) for item in candidate_units]))

        if current_units and candidate_length > chunk_size:
            chunks.append(_build_chunk_payload(current_units))
            current_units = _fit_overlap_units(
                _collect_overlap_units(current_units, chunk_overlap),
                next_unit=unit,
                chunk_size=chunk_size,
            )

        current_units.append(unit)

    if current_units:
        chunks.append(_build_chunk_payload(current_units))

    return chunks


def _build_chunk_payload(units: list[Block]) -> dict:
    return {
        "content": _join_blocks([str(unit.get("text", "")) for unit in units]),
        "paragraph_index_start": _range_start(units, "paragraph_index"),
        "paragraph_index_end": _range_end(units, "paragraph_index"),
        "block_index_start": _range_start(units, "block_index"),
        "block_index_end": _range_end(units, "block_index"),
    }


def _normalize_block(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _join_blocks(blocks: list[str]) -> str:
    return "\n\n".join(block for block in blocks if block).strip()


def _split_long_block(block: Block, chunk_size: int) -> list[Block]:
    text = str(block.get("text", ""))
    sentences = [
        sentence.strip()
        for sentence in re.split(r"(?<=[。！？!?；;])\s+|(?<=[。！？!?；;])", text)
        if sentence.strip()
    ]

    if len(sentences) <= 1:
        return [
            {
                **block,
                "text": text[index : index + chunk_size].strip(),
            }
            for index in range(0, len(text), chunk_size)
            if text[index : index + chunk_size].strip()
        ]

    pieces: list[Block] = []
    current = ""
    for sentence in sentences:
        candidate = f"{current}{sentence}" if current else sentence
        if current and len(candidate) > chunk_size:
            pieces.append({**block, "text": current.strip()})
            current = sentence
        else:
            current = candidate

    if current.strip():
        pieces.append({**block, "text": current.strip()})

    return pieces


def _collect_overlap_units(units: list[Block], overlap_chars: int) -> list[Block]:
    if overlap_chars <= 0:
        return []

    overlap: list[Block] = []
    total = 0
    for unit in reversed(units):
        overlap.insert(0, unit)
        total += len(str(unit.get("text", "")))
        if total >= overlap_chars:
            break

    return overlap


def _fit_overlap_units(
    overlap_units: list[Block],
    next_unit: Block,
    chunk_size: int,
) -> list[Block]:
    fitted = list(overlap_units)
    next_text = str(next_unit.get("text", ""))

    while fitted:
        candidate_texts = [str(unit.get("text", "")) for unit in fitted] + [next_text]
        if len(_join_blocks(candidate_texts)) <= chunk_size:
            break
        fitted.pop(0)

    return fitted


def _find_meaningful_paragraph(sequence: list[dict], index: int, direction: int) -> dict | None:
    cursor = index + direction
    while 0 <= cursor < len(sequence):
        item = sequence[cursor]
        if item.get("type") == "paragraph" and _is_meaningful_anchor_text(str(item.get("text", ""))):
            return item
        cursor += direction
    return None


def _is_meaningful_anchor_text(text: str) -> bool:
    normalized = _normalize_block(text)
    if len(normalized) < 4:
        return False
    meaningless = {"如下图", "见下图", "如下", "见图", "如图", "如下所示", "见下页"}
    return normalized not in meaningless


def _build_anchor_text(anchor_before: dict | None, anchor_after: dict | None) -> str:
    texts = []
    if anchor_before is not None:
        texts.append(str(anchor_before.get("text", "")).strip())
    if anchor_after is not None:
        after_text = str(anchor_after.get("text", "")).strip()
        if after_text and after_text not in texts:
            texts.append(after_text)
    return "\n".join(text for text in texts if text)


def _find_nearest_pdf_blocks(
    text_blocks: list[Block],
    image_bbox,
) -> tuple[Block | None, Block | None]:
    if image_bbox is None:
        return None, None

    image_center_y = (float(image_bbox.y0) + float(image_bbox.y1)) / 2
    before_candidates = []
    after_candidates = []

    for block in text_blocks:
        bbox = block.get("bbox")
        if not isinstance(bbox, tuple) or len(bbox) != 4:
            continue

        block_center_y = (float(bbox[1]) + float(bbox[3])) / 2
        delta = abs(block_center_y - image_center_y)
        if block_center_y <= image_center_y:
            before_candidates.append((delta, block))
        else:
            after_candidates.append((delta, block))

    before_candidates.sort(key=lambda item: item[0])
    after_candidates.sort(key=lambda item: item[0])
    before = before_candidates[0][1] if before_candidates else None
    after = after_candidates[0][1] if after_candidates else None
    return before, after


def _range_start(items: list[Block], key: str) -> int:
    values = [int(item.get(key, 0) or 0) for item in items if int(item.get(key, 0) or 0) > 0]
    return min(values) if values else 0


def _range_end(items: list[Block], key: str) -> int:
    values = [int(item.get(key, 0) or 0) for item in items if int(item.get(key, 0) or 0) > 0]
    return max(values) if values else 0


def _guess_content_type(name_or_ext: str) -> str:
    ext = os.path.splitext(name_or_ext)[1].lower() if "." in name_or_ext else f'.{name_or_ext.lower()}'
    if ext in {".jpg", ".jpeg"}:
        return "image/jpeg"
    if ext == ".gif":
        return "image/gif"
    if ext == ".webp":
        return "image/webp"
    if ext == ".bmp":
        return "image/bmp"
    return "image/png"


def _is_docx_heading(paragraph) -> bool:
    style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
    if _looks_like_heading(paragraph.text):
        return True
    return bool(re.match(r"^(heading|标题)\s*\d*$", style_name.strip(), flags=re.IGNORECASE))


def _get_docx_heading_level(paragraph) -> int:
    style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
    match = re.search(r"(\d+)", style_name)
    if match:
        return max(1, min(int(match.group(1)), 6))
    return _infer_heading_level(paragraph.text)


def _looks_like_heading(text: str) -> bool:
    normalized = _normalize_block(text)
    if not normalized or len(normalized) > TITLE_MAX_LENGTH:
        return False
    if normalized.endswith(("。", "；", ";", ":", "：")):
        return False
    patterns = (
        r"^#{1,6}\s+.+$",
        r"^第[一二三四五六七八九十百0-9]+[章节部分篇]\s*.*$",
        r"^[一二三四五六七八九十]+[、.]\s*.+$",
        r"^[0-9]+(?:\.[0-9]+){0,3}\s+.+$",
        r"^[0-9]+[、.)]\s*.+$",
        r"^[（(][0-9一二三四五六七八九十]+[）)]\s*.+$",
    )
    return any(re.match(pattern, normalized) for pattern in patterns)


def _infer_heading_level(text: str) -> int:
    normalized = _normalize_block(text)
    if re.match(r"^#{1,6}\s+.+$", normalized):
        return min(len(normalized) - len(normalized.lstrip("#")), 6)
    if re.match(r"^第[一二三四五六七八九十百0-9]+[章节部分篇]", normalized):
        return 1
    if re.match(r"^[一二三四五六七八九十]+[、.]", normalized):
        return 2
    if re.match(r"^[0-9]+\.[0-9]+\.[0-9]+", normalized):
        return 4
    if re.match(r"^[0-9]+\.[0-9]+", normalized):
        return 3
    if re.match(r"^[0-9]+[、.)]", normalized):
        return 3
    if re.match(r"^[（(][0-9一二三四五六七八九十]+[）)]", normalized):
        return 4
    return 2


def _update_heading_stack(stack: list[str], heading: str, level: int) -> list[str]:
    normalized_heading = _normalize_block(heading)
    next_stack = list(stack[: max(level - 1, 0)])
    next_stack.append(normalized_heading)
    return next_stack


def _load_pymupdf():
    for module_name in ("fitz", "pymupdf"):
        try:
            return import_module(module_name)
        except ModuleNotFoundError:
            continue

    raise RuntimeError(
        "PDF 解析依赖未安装，请在 backend 环境中执行 `pip install pymupdf` "
        "或重新同步 `pyproject.toml` 依赖后再试。"
    )
